"""Genera un reporte de ventas en Excel a partir de un CSV."""
import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNAS = ["fecha", "producto", "cantidad", "precio"]
FORMATO_MONEDA = '"$"#,##0'


class ErrorDatos(Exception):
    """Problema con el archivo de entrada, con mensaje legible."""


def leer_datos(entrada):
    ruta = Path(entrada)
    if not ruta.is_file():
        raise ErrorDatos(f"No existe el archivo: {ruta}")
    try:
        with open(ruta, encoding="utf-8-sig", errors="replace") as f:
            primera = f.readline()
        if not primera.strip():
            raise ErrorDatos("El archivo está vacío.")
        sep = ";" if primera.count(";") > primera.count(",") else ","
        df = pd.read_csv(ruta, encoding="utf-8-sig", sep=sep,
                         decimal="," if sep == ";" else ".")
    except ErrorDatos:
        raise
    except pd.errors.EmptyDataError:
        raise ErrorDatos("El archivo está vacío.")
    except Exception as e:
        raise ErrorDatos(f"No se pudo leer el CSV: {e}")

    df.columns = [str(c).strip().lower() for c in df.columns]
    faltan = [c for c in COLUMNAS if c not in df.columns]
    if faltan:
        raise ErrorDatos(f"Faltan columnas: {', '.join(faltan)}. "
                         f"Se esperaba: {', '.join(COLUMNAS)}.")
    df = df[COLUMNAS].dropna(how="all").copy()
    if df.empty:
        raise ErrorDatos("El archivo no tiene filas de datos.")

    texto = df["fecha"].astype("string").str.strip()
    iso = pd.to_datetime(texto, format="%Y-%m-%d", errors="coerce")
    latina = pd.to_datetime(texto, format="%d/%m/%Y", errors="coerce")
    df["fecha"] = iso.fillna(latina)
    for col in ("cantidad", "precio"):
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df["producto"] = df["producto"].astype("string").str.strip()

    malas = df[df.isna().any(axis=1)]
    if not malas.empty:
        filas = ", ".join(str(i + 2) for i in malas.index[:10])
        raise ErrorDatos(f"Hay {len(malas)} fila(s) con datos vacíos o "
                         f"inválidos (filas {filas}). Corrígelas y vuelve a ejecutar.")
    if (df[["cantidad", "precio"]] < 0).any().any():
        raise ErrorDatos("Hay cantidades o precios negativos.")

    df["total"] = df["cantidad"] * df["precio"]
    return df


def dar_formato(hoja, columnas_moneda=(), columnas_fecha=()):
    encabezado = PatternFill("solid", fgColor="1F4E78")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = encabezado
        celda.alignment = Alignment(horizontal="center")
    for idx in columnas_moneda:
        for fila in hoja.iter_rows(min_row=2, min_col=idx, max_col=idx):
            fila[0].number_format = FORMATO_MONEDA
    for idx in columnas_fecha:
        for fila in hoja.iter_rows(min_row=2, min_col=idx, max_col=idx):
            fila[0].number_format = "yyyy-mm-dd"
    for col in hoja.columns:
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        hoja.column_dimensions[get_column_letter(col[0].column)].width = ancho + 4
    hoja.freeze_panes = "A2"


def generar_reporte(entrada, salida):
    df = leer_datos(entrada)
    resumen = (df.groupby("producto")
                 .agg(unidades=("cantidad", "sum"), total=("total", "sum"))
                 .sort_values("total", ascending=False).reset_index())
    resumen["participacion"] = resumen["total"] / resumen["total"].sum()
    fila_total = pd.DataFrame([{"producto": "TOTAL",
                                "unidades": resumen["unidades"].sum(),
                                "total": resumen["total"].sum(),
                                "participacion": 1.0}])
    resumen = pd.concat([resumen, fila_total], ignore_index=True)

    salida = Path(salida)
    if salida.suffix.lower() != ".xlsx":
        salida = salida.with_suffix(".xlsx")
    try:
        with pd.ExcelWriter(salida, engine="openpyxl") as xl:
            df.sort_values("fecha").to_excel(xl, sheet_name="Detalle", index=False)
            resumen.to_excel(xl, sheet_name="Por producto", index=False)
            dar_formato(xl.sheets["Detalle"], columnas_moneda=(4, 5), columnas_fecha=(1,))
            hoja = xl.sheets["Por producto"]
            dar_formato(hoja, columnas_moneda=(3,))
            for fila in hoja.iter_rows(min_row=2, min_col=4, max_col=4):
                fila[0].number_format = "0.0%"
            for celda in hoja[hoja.max_row]:
                celda.font = Font(bold=True)
    except PermissionError:
        raise ErrorDatos(f"No se puede escribir {salida}. ¿Está abierto en otro programa?")
    return salida


def main():
    p = argparse.ArgumentParser(description="Genera un reporte de ventas en Excel desde un CSV.")
    p.add_argument("entrada", help="CSV con columnas: " + ", ".join(COLUMNAS))
    p.add_argument("salida", nargs="?", default="reporte.xlsx")
    args = p.parse_args()
    try:
        ruta = generar_reporte(args.entrada, args.salida)
    except ErrorDatos as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(1)
    print(f"Reporte listo: {ruta}")


if __name__ == "__main__":
    main()
