import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from generar_reporte import ErrorDatos, generar_reporte, leer_datos  # noqa: E402

ENCABEZADO = "fecha,producto,cantidad,precio\n"


def csv(tmp_path, contenido, nombre="datos.csv", binario=False):
    ruta = tmp_path / nombre
    if binario:
        ruta.write_bytes(contenido)
    else:
        ruta.write_text(contenido, encoding="utf-8")
    return ruta


def test_caso_feliz(tmp_path):
    entrada = csv(tmp_path, ENCABEZADO + "2026-10-01,Arroz,10,4500\n2026-10-02,Aceite,5,12000\n")
    salida = generar_reporte(entrada, tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(salida)
    assert wb.sheetnames == ["Detalle", "Por producto"]
    filas = list(wb["Por producto"].iter_rows(values_only=True))
    assert filas[-1][0] == "TOTAL"
    assert filas[-1][2] == 10 * 4500 + 5 * 12000


def test_archivo_inexistente(tmp_path):
    with pytest.raises(ErrorDatos, match="No existe"):
        leer_datos(tmp_path / "nada.csv")


def test_archivo_vacio(tmp_path):
    with pytest.raises(ErrorDatos, match="vacío"):
        leer_datos(csv(tmp_path, ""))


def test_solo_encabezado(tmp_path):
    with pytest.raises(ErrorDatos, match="no tiene filas"):
        leer_datos(csv(tmp_path, ENCABEZADO))


def test_columnas_faltantes(tmp_path):
    with pytest.raises(ErrorDatos, match="Faltan columnas"):
        leer_datos(csv(tmp_path, "fecha,producto\n2026-10-01,Arroz\n"))


def test_datos_invalidos_indica_fila(tmp_path):
    with pytest.raises(ErrorDatos, match="filas 3"):
        leer_datos(csv(tmp_path, ENCABEZADO + "2026-10-01,Arroz,1,4500\nhoy,Arroz,x,4500\n"))


def test_negativos(tmp_path):
    with pytest.raises(ErrorDatos, match="negativos"):
        leer_datos(csv(tmp_path, ENCABEZADO + "2026-10-01,Arroz,-1,4500\n"))


def test_punto_y_coma_decimal_coma_y_bom(tmp_path):
    contenido = "\ufeffFecha;Producto;Cantidad;Precio\n2026-10-01; Arroz ;2,5;4000\n".encode("utf-8")
    df = leer_datos(csv(tmp_path, contenido, binario=True))
    assert df.loc[0, "cantidad"] == 2.5
    assert df.loc[0, "producto"] == "Arroz"
    assert df.loc[0, "total"] == 10000


def test_fecha_dia_mes_anio(tmp_path):
    df = leer_datos(csv(tmp_path, ENCABEZADO + "03/10/2026,Arroz,1,4500\n"))
    assert (df.loc[0, "fecha"].month, df.loc[0, "fecha"].day) == (10, 3)


def test_extension_se_corrige(tmp_path):
    entrada = csv(tmp_path, ENCABEZADO + "2026-10-01,Arroz,1,4500\n")
    salida = generar_reporte(entrada, tmp_path / "sin_extension")
    assert salida.suffix == ".xlsx" and salida.exists()
